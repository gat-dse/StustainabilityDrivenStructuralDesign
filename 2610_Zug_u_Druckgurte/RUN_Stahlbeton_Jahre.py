"""
Stahlbeton-Gurte im Verlauf der Jahre (1990 - 2050)

Last-Verformungs-Verhalten N(u) eines Stahlbeton-Gurts (C30/37, L0 = 1 m) bei einem fixen
Budget von 100 kgCO2 pro m Gurtlänge, für verschiedene Betone aus dem treeze-Betonrechner
(Excel "Zementemissionen-jh.xlsx").

  Zug:  Zuggurtmodell mit Tension Stiffening (N_rc aus RUN.py)
  Druck: Beton (Sargin-Parabel, eps_c2d = 3.5 ‰) + Betonstahl

Die Materialmodelle werden aus RUN.py übernommen; hier ändern sich nur die CO2-Werte:
  e_c  [kgCO2/m3 Beton]  = "kgCO2 total" (Zeile 12) der jeweiligen Betonsorte
  e_s  [kgCO2/m3 Stahl]  = kgCO2 pro 100 kg Bewehrung (Zeilen 14/16) / 100 kg * 7850 kg/m3
                           Stahlsorte wählbar über STEEL ("CH mix" oder "Gerlafingen")
"""

from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle

import RUN as base   # Materialmodelle, Plot-Stil und Hilfsfunktionen

mm, MPa = base.mm, base.MPa
HERE = Path(__file__).resolve().parent
XLSX = HERE / "Zementemissionen-jh.xlsx"

STEEL = "CH mix"        # Betonstahl: "CH mix" (Excel Zeile 14) oder "Gerlafingen" (Zeile 16)
RHO = base.RHO          # geometrischer Bewehrungsgehalt (wie RUN.py, 5 %)
GAMMA_S = 7850          # kg/m3

# Betonsorten: (Spalte im Excel, Jahr, Kurzname, Farbe)
CONCRETES = [
    ("D", 1990, "CEM I (fossil)",      "#5b3a29"),
    ("E", 2026, "CEM I (CH mix)",      "#b5651d"),
    ("J", 2026, "ND Kibeco (Reco200)", "#2e9e44"),
    ("K", 2050, "Beton 2050",          "#1f78b4"),
    ("L", 2050, "Beton 2050 mit CCUS", "#6a3d9a"),
]


# =============================================================================
# 1) CO2-Werte aus dem Excel
# =============================================================================
def read_emissions(path=XLSX):
    """Liest e_c (Zeile 12) sowie die Stahlwerte (Zeilen 14/16) aus dem Excel.
    Excel speichert berechnete Formelwerte; die Datei muss einmal in Excel gespeichert sein."""
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).active
    steel = {"Gerlafingen": ws["D16"].value, "CH mix": ws["D14"].value}   # kgCO2 / 100 kg
    rows = []
    for col, year, name, color in CONCRETES:
        e_c = ws[f"{col}12"].value
        if e_c is None:
            raise ValueError(f"Zelle {col}12 leer – Excel einmal speichern (Formelwerte fehlen).")
        rows.append(dict(col=col, year=year, name=name, color=color, e_c=float(e_c)))
    e_s = {k: v / 100 * GAMMA_S for k, v in steel.items()}                # kgCO2 / m3 Stahl
    return rows, e_s


def N_rc_e(u, e_c, e_s):
    """Stahlbeton-Gurt N(u) [N] für die CO2-Werte e_c (Beton) und e_s (Stahl)."""
    return base.N_rc(u, RHO, e_c, e_s)


# =============================================================================
# 2) Plot N(u) + Querschnitte
# =============================================================================
def plot(rows, e_s, u_min=-4, u_max=4, n=4001, fname="stahlbeton_jahre.png"):
    uu = np.linspace(u_min, u_max, n) * mm
    es = e_s[STEEL]

    fig = plt.figure(figsize=(15, 13))
    gs = fig.add_gridspec(2, 1, height_ratios=[3.0, 1.0], hspace=0.12,
                          left=0.08, right=0.90, top=0.97, bottom=0.10)
    ax = fig.add_subplot(gs[0])
    axc = fig.add_subplot(gs[1])

    tens, comp = [], []
    for r in rows:
        col = r["color"]
        e_c = r["e_c"]
        N = N_rc_e(uu, e_c, es) / 1e6
        ax.plot(uu / mm, base._hide_zero(uu, N), color=col, lw=2.8, zorder=4,
                label=f"{r['year']}: {r['name']}  ($e_c$ = {r['e_c']:.0f} kgCO$_2$/m$^3$)")
        for sign, store, x_edge in ((+1, tens, u_max), (-1, comp, u_min)):
            u_ult = base._ultimate_u(lambda u, k: N_rc_e(u, e_c, es), None, sign)
            if abs(u_ult) > abs(x_edge) * mm:
                store.append((float(N_rc_e(x_edge * mm, e_c, es) / 1e6), u_ult, col))
    if tens:
        base._end_arrows(ax, tens, u_max, +1, min_gap=0.45)
    if comp:
        base._end_arrows(ax, comp, u_min, -1, min_gap=0.45)

    # Achsen im Ursprung
    Nall = [l.get_ydata() for l in ax.lines]
    ymax = np.nanmax([np.nanmax(y) for y in Nall]) * 1.15
    ymin = np.nanmin([np.nanmin(y) for y in Nall]) * 1.15
    ax.set_xlim(u_min - 0.1, u_max + 0.3)
    ax.set_ylim(ymin, ymax)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_position(("data", 0))
    ax.spines["bottom"].set_position(("data", 0))
    ax.plot(1, 0, ">k", ms=11, transform=ax.get_yaxis_transform(), clip_on=False)
    ax.plot(0, 1, "^k", ms=11, transform=ax.get_xaxis_transform(), clip_on=False)
    ax.set_xticks(np.arange(u_min, u_max + 0.01, 0.5))
    ax.set_yticks(np.arange(np.ceil(ymin / 2.5) * 2.5, ymax, 2.5))
    ax.xaxis.set_major_formatter(lambda x, p: "" if abs(x) < 1e-9 else f"{x:.1f}")
    ax.yaxis.set_major_formatter(lambda y, p: "" if abs(y) < 1e-9 else f"{y:g}")
    ax.tick_params(direction="in", length=5, labelsize=13)
    for lab in ax.get_xticklabels() + ax.get_yticklabels():
        lab.set_bbox(dict(facecolor="white", edgecolor="none", alpha=0.7, pad=0.5))
    ax.text(0.15, ymax, "$N_{Rd}$ [MN]", fontsize=17, ha="left", va="top")
    ax.text(u_max + 0.35, -0.03 * (ymax - ymin), "$u$ [mm]", fontsize=17, ha="left", va="top",
            clip_on=False)
    box = dict(boxstyle="square,pad=0.3", fc="white", ec="k", lw=1.2)
    ax.text(1.6, ymax * 0.97, "Zuggurt", fontsize=18, ha="center", va="top", bbox=box)
    ax.text(-2.6, ymin * 0.97, "Druckgurt", fontsize=18, ha="center", va="bottom", bbox=box)

    leg1 = ax.legend(loc="upper left", frameon=False, fontsize=12.5,
                     title=f"Stahlbeton C30/37, $\\rho_s$ = {RHO:.0%}", title_fontsize=13.5)
    leg1._legend_box.align = "left"
    ax.add_artist(leg1)
    ax.text(0.98, 0.04, f"Betonstahl {STEEL}: $e_s$ = {base._fmt(es)} kgCO$_2$/m$^3$\n"
            "Budget: 100 kgCO$_2$ pro m Gurtlänge", transform=ax.transAxes,
            ha="right", va="bottom", fontsize=13, linespacing=1.6)

    _sections(axc, rows, es)
    ax.text(-0.06, 1.0, "(a)", transform=ax.transAxes, fontsize=17, va="top")
    axc.text(-0.06, 1.0, "(b)", transform=axc.transAxes, fontsize=17, va="top")
    fig.savefig(HERE / fname, dpi=150, bbox_inches="tight", pad_inches=0.2)
    return fig


def _sections(axc, rows, e_s):
    """Querschnitte massstäblich als Quadrate."""
    gk_c, gk_s = base.gamma["c"] * base.g / 1000, base.gamma["s"] * base.g / 1000
    W = 1150.0
    for i, r in enumerate(rows):
        A = base.A_rc(RHO, r["e_c"], e_s)
        a = np.sqrt(A) / mm
        x0 = i * W
        axc.add_patch(Rectangle((x0, 0), a, a, color=r["color"], alpha=0.7, lw=0))
        txt = ["$\\bf{" + str(r["year"]) + "}$", r["name"],
               f"$e_c$ = {r['e_c']:.0f} kgCO$_2$/m$^3$",
               f"$A$ = {base._fmt(A * 1e6)} mm$^2$", f"$a$ = {base._fmt(a)} mm",
               f"$g_{{0k}}$ = {A * ((1 - RHO) * gk_c + RHO * gk_s):.1f} kN/m"]
        axc.annotate("\n".join(txt), xy=(x0, 0), xytext=(0, -10), textcoords="offset points",
                     ha="left", va="top", fontsize=12.5, linespacing=1.5, annotation_clip=False)
    axc.set_xlim(-30, max(len(rows), 6) * W)
    axc.set_ylim(0, 1000)
    axc.set_aspect("equal", anchor="NW")
    axc.axis("off")


# =============================================================================
# 3) Tabelle
# =============================================================================
def table(rows, e_s):
    out = []
    for r in rows:
        line = {"Jahr": r["year"], "Beton": r["name"], "e_c [kgCO2/m3]": r["e_c"]}
        for steel, es in [(STEEL, e_s[STEEL])]:
            A = base.A_rc(RHO, r["e_c"], es)
            uu = np.linspace(-3.5, 60, 6000) * mm
            N = base.N_rc(uu, RHO, r["e_c"], es) / 1e6
            line[f"A {steel} [mm²]"] = A * 1e6
            line[f"a {steel} [mm]"] = np.sqrt(A) / mm
            line[f"N_Zug,max {steel} [MN]"] = N.max()
            line[f"N_Druck,max {steel} [MN]"] = -N.min()
        out.append(line)
    import pandas as pd
    return pd.DataFrame(out).set_index(["Jahr", "Beton"])


if __name__ == "__main__":
    rows, e_s = read_emissions()
    tab = table(rows, e_s)
    import pandas as pd
    with pd.option_context("display.float_format", "{:,.1f}".format, "display.width", 250,
                           "display.max_columns", None):
        print(tab)
    tab.to_csv(HERE / "stahlbeton_jahre.csv", float_format="%.2f")
    plot(rows, e_s)
    plt.show()
